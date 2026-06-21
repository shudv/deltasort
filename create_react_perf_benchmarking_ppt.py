from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(9, 12, 20)
BG_ALT = RGBColor(15, 20, 32)
TEXT = RGBColor(240, 244, 255)
MUTED = RGBColor(151, 167, 196)
ACCENT = RGBColor(0, 226, 204)
ACCENT_2 = RGBColor(255, 125, 61)
DANGER = RGBColor(255, 83, 112)


def add_background(slide, motif_variant=0):
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    base.fill.solid()
    base.fill.fore_color.rgb = BG
    base.line.fill.background()

    haze = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(0.2), Inches(12.9), Inches(7.1))
    haze.fill.solid()
    haze.fill.fore_color.rgb = BG_ALT
    haze.fill.transparency = 0.45
    haze.line.fill.background()

    # Recurring motif: profiler-like skyline near bottom edge.
    start_x = 0.6
    bar_w = 0.22
    heights = [0.3, 0.5, 0.9, 0.35, 1.4, 0.7, 0.45, 1.8, 0.6, 1.2, 0.4, 0.95, 0.5, 1.5, 0.65, 1.1]
    shift = motif_variant % 4
    for i, h in enumerate(heights):
        x = Inches(start_x + i * (bar_w + 0.06))
        height = Inches(h + (0.08 * ((i + shift) % 3)))
        y = Inches(7.2) - height
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(bar_w), height)
        bar.fill.solid()
        if i % 2 == 0:
            bar.fill.fore_color.rgb = ACCENT
        else:
            bar.fill.fore_color.rgb = ACCENT_2
        bar.fill.transparency = 0.72
        bar.line.fill.background()


def add_title(slide, text, top=0.6, size=52, color=TEXT):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.8), Inches(1.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Avenir Next"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT


def add_body(slide, text, top=2.15, size=26, color=MUTED):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(2.4))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Avenir Next"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT


def add_section_tag(slide, text):
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.2), Inches(3.7), Inches(0.55))
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor(26, 37, 61)
    tag.fill.transparency = 0.1
    tag.line.fill.background()
    tf = tag.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.runs[0].font.name = "Avenir Next"
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER


def add_myth_reality_split(slide, myth, reality):
    left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.15), Inches(5.95), Inches(2.75))
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(54, 24, 42)
    left.fill.transparency = 0.25
    left.line.fill.background()

    right = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.45), Inches(3.15), Inches(5.95), Inches(2.75))
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(20, 58, 68)
    right.fill.transparency = 0.2
    right.line.fill.background()

    myth_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.45), Inches(5.2), Inches(1.8))
    mtf = myth_box.text_frame
    mtf.clear()
    p = mtf.paragraphs[0]
    p.text = "MYTH"
    p.runs[0].font.name = "Avenir Next"
    p.runs[0].font.size = Pt(16)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = DANGER
    p2 = mtf.add_paragraph()
    p2.text = myth
    p2.runs[0].font.name = "Avenir Next"
    p2.runs[0].font.size = Pt(24)
    p2.runs[0].font.bold = True
    p2.runs[0].font.color.rgb = TEXT

    real_box = slide.shapes.add_textbox(Inches(6.75), Inches(3.45), Inches(5.2), Inches(1.8))
    rtf = real_box.text_frame
    rtf.clear()
    p = rtf.paragraphs[0]
    p.text = "REALITY"
    p.runs[0].font.name = "Avenir Next"
    p.runs[0].font.size = Pt(16)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = ACCENT
    p2 = rtf.add_paragraph()
    p2.text = reality
    p2.runs[0].font.name = "Avenir Next"
    p2.runs[0].font.size = Pt(24)
    p2.runs[0].font.bold = True
    p2.runs[0].font.color.rgb = TEXT


def add_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes


def build_deck(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # Slide 1: provocative opener
    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=0)
    add_title(s, "Your fastest React optimization probably made it slower.")
    add_body(s, "Memoization can be a tax, not a discount.", top=2.25, size=30, color=ACCENT_2)
    add_notes(
        s,
        "Open with surprise. Senior engineers expect nuance, but not this punchline. "
        "Frame the talk around avoiding confident mistakes in browser performance work. "
        "Tease that every claim today comes from profiles and production incidents, not folklore.",
    )

    # Part 1
    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=1)
    add_section_tag(s, "PART 1 - Performance is Tricky")
    add_title(s, "Memoization is not free.", top=1.0)
    add_myth_reality_split(s, "Cache everything.", "Cache only where compute beats cache overhead.")
    add_notes(
        s,
        "Explain hidden costs: dependency tracking, extra memory churn, stale cache invalidation, and referential equality gymnastics. "
        "Mention that low-cost calculations become slower when wrapped with caching ceremony.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=2)
    add_section_tag(s, "PART 1 - Performance is Tricky")
    add_title(s, "Your runtime cost model is usually wrong.", top=1.0)
    add_body(s, "CPU, GC, layout, paint, network: different bottlenecks, different truths.", top=2.4, size=28)
    add_notes(
        s,
        "Challenge simplistic big-O thinking in UI code. "
        "In browsers, one avoidable reflow can dominate dozens of JavaScript operations. "
        "A correct model includes scheduler behavior, main-thread contention, and render pipeline phases.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=3)
    add_section_tag(s, "PART 1 - Performance is Tricky")
    add_title(s, "Svelte says \"fast\". Where are the profiles?", top=1.0)
    add_body(s, "Marketing velocity without profiling evidence should trigger engineering skepticism.", top=2.45, size=28, color=ACCENT)
    add_notes(
        s,
        "Do not bash Svelte itself. The point is broader: performance claims without reproducible profiles are weak evidence. "
        "Ask audience to demand traces, test conditions, and user-centric metrics before accepting speed narratives.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=0)
    add_section_tag(s, "PART 1 - Performance is Tricky")
    add_title(s, "Baselines orient you. Baselines can also trap you.", top=1.0)
    add_body(s, "Great compass. Bad autopilot.", top=2.55, size=34, color=ACCENT_2)
    add_notes(
        s,
        "Baseline benchmarks are essential for sanity checks and regression detection. "
        "But blindly optimizing to synthetic baselines can overfit away from real user paths. "
        "Use baselines as map, not mandate.",
    )

    # Part 2
    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=1)
    add_section_tag(s, "PART 2 - Getting Performance Right")
    add_title(s, "Principle 1: Model the true cost.", top=1.0)
    add_body(s, "Know what is expensive here, now, in this browser path.", top=2.35, size=28)
    add_notes(
        s,
        "Encourage explicit hypotheses tied to measurable components: script, style, layout, paint, compositing, and network waits. "
        "Different pages and devices shift bottlenecks.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=2)
    add_section_tag(s, "PART 2 - Getting Performance Right")
    add_title(s, "Principle 2: Profile before optimizing.", top=1.0)
    add_body(s, "No profile, no optimization. Only guesses.", top=2.55, size=34, color=ACCENT)
    add_notes(
        s,
        "Set a strict team norm: optimization PRs should include before and after traces. "
        "Call out where time moved, not just total time. "
        "Profiles often show the opposite of first intuition.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=3)
    add_section_tag(s, "PART 2 - Getting Performance Right")
    add_title(s, "Principle 3: Users feel latency, not your benchmark score.", top=1.0)
    add_body(s, "Perceived performance is product performance.", top=2.55, size=32, color=ACCENT_2)
    add_notes(
        s,
        "Tie engineering metrics to user perception: interaction latency, jank, and visual stability. "
        "A small metric win can be invisible; a tiny visual hitch can destroy confidence.",
    )

    # Part 3
    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=0)
    add_section_tag(s, "PART 3 - Real Production Examples")
    add_title(s, "Case Study A", top=0.95, size=48)
    add_myth_reality_split(s, "React diffing is too slow.", "Forced reflow loop in a custom tooltip hook.")
    add_body(s, "Symptom: typing lag in dense tables.", top=2.2, size=27, color=MUTED)
    add_notes(
        s,
        "Symptom: users reported input lag in a grid with hover tooltips. "
        "Misleading first hypothesis blamed React reconciliation due to many cells. "
        "Profile showed repeated read-after-write layout invalidation in a hook measuring tooltip bounds on every keypress. "
        "Fix: batch reads and writes in requestAnimationFrame and decouple tooltip updates from typing path.",
    )

    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=1)
    add_section_tag(s, "PART 3 - Real Production Examples")
    add_title(s, "Case Study B", top=0.95, size=48)
    add_myth_reality_split(s, "Virtualization library bug.", "Unstable keys remounted rows, nuking memoization.")
    add_body(s, "Symptom: scroll jank after filter changes.", top=2.2, size=27, color=MUTED)
    add_notes(
        s,
        "Symptom: smooth scroll degraded badly only after filtering results. "
        "First hypothesis targeted the virtualization library internals. "
        "Root cause: key generation used index plus transient filter state, causing frequent remounts and cache invalidation. "
        "Fix: stable entity IDs as keys and memo boundaries around row renderers.",
    )

    # Part 4
    s = prs.slides.add_slide(blank)
    add_background(s, motif_variant=2)
    add_section_tag(s, "PART 4 - Recap")
    add_title(s, "Carry these into your next perf incident:", top=0.95, size=46)
    add_body(s, "Model reality. Profile first. Optimize for felt speed.", top=2.45, size=34, color=ACCENT)
    add_notes(
        s,
        "Close by restating the practical discipline: build accurate cost models, let profiles arbitrate, and optimize what users perceive. "
        "Invite the team to standardize perf postmortems around these three checkpoints.",
    )

    prs.save(output_path)


if __name__ == "__main__":
    build_deck("/Users/shubham/Repos/deltasort/React_Webapp_Browser_Performance_Benchmarking.pptx")
